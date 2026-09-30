"""AXIS Sustainability Readiness Review.

Statuses, counts, priorities, evidence IDs and registers are computed from the saved audit.
Only narrative text (letter summary, gap wording, action plan) comes from the optional AI draft,
and it is validated against the audit before use.
"""
import io
import re
from datetime import datetime, timezone

PILLARS = [("A", "Sustainable management"), ("B", "Socioeconomic impacts"), ("C", "Cultural impacts"), ("D", "Environmental impacts")]
OTHER = ("X", "Additional requirements")
STATUSES = ["Met", "Partly met", "Not met", "Not evidenced", "Not sampled", "Not applicable (provisional)"]
NEEDS_WORK = {"Partly met", "Not met", "Not evidenced"}
EVIDENCE_PREFIX = {"document": "D", "record": "D", "interview": "I", "observation": "O", "photo": "P"}
EVIDENCE_LABEL = {"document": "Document", "record": "Record", "interview": "Interview", "observation": "Observation", "photo": "Photo"}
EVIDENCE_TYPES = ["observed", "documented", "staff-reported", "unverified"]
PLACEHOLDER = "[Reviewer to complete]"

PROFILE_FIELDS = {
    "hotel_name": 255, "location": 255, "hotel_contact_name": 255, "hotel_contact_role": 255,
    "purpose": 500, "standard_edition": 100, "standard_used": 500, "review_date": 100, "reviewer": 255,
    "reviewer_contact": 255, "report_reference": 100, "report_date": 100, "confidentiality": 1000,
    "staff_present": 1000, "review_type": 255, "duration": 255, "method": 1000,
    "hotel_type": 255, "rooms": 255, "staff": 255, "facilities": 1000, "occupancy": 255,
    "water_sources": 255, "wastewater": 255, "existing_certifications": 500,
    "people_interviewed": 4000, "areas_inspected": 4000, "areas_not_inspected": 2000,
    "plan_changes": 4000, "key_figures": 6000, "legal_items": 8000,
    "prepared_by": 255, "reviewed_by": 255, "issued_to": 255, "next_step": 500,
}


def default_profile(bundle):
    audit = bundle["audit"]
    template = bundle.get("template") or {}
    edition = "v4.0" if template.get("id") == "gstc-hotel-v4" else template.get("version", "")
    return {
        "hotel_name": audit.get("siteName", ""),
        "purpose": "Identify what the hotel must address to achieve GSTC certification",
        "standard_edition": edition,
        "standard_used": f"{template.get('title', 'GSTC Hotel Standard')} {template.get('version', '')}".strip(),
        "reviewer": audit.get("leadAuditor", {}).get("name", ""),
        "review_type": "One-day on-site readiness review",
        "duration": "One reviewer-day on site, plus document review before and after",
        "method": "Pre-visit questionnaire and document request; document review; site walk-through; staff interviews; closing debrief",
        "confidentiality": "Confidential to the hotel. Anonymised data may be pooled for benchmarking under the signed agreement.",
        "prepared_by": audit.get("leadAuditor", {}).get("name", ""),
    }


def merge_profile(bundle, saved):
    profile = default_profile(bundle)
    for key, value in (saved or {}).items():
        if key in PROFILE_FIELDS and isinstance(value, str) and value.strip():
            profile[key] = value.strip()[:PROFILE_FIELDS[key]]
    return profile


def rows(text, columns):
    """Parse 'a | b | c' lines typed by the reviewer into fixed-width table rows."""
    result = []
    for line in (text or "").splitlines():
        if not line.strip():
            continue
        cells = [cell.strip() for cell in line.split("|")][:columns]
        result.append(cells + [""] * (columns - len(cells)))
    return result


def lines(text):
    return [line.strip().lstrip("-•* ").strip() for line in (text or "").splitlines() if line.strip()]


def _clause_key(code):
    match = re.match(r"^([A-Z])(\d+)$", code or "")
    return (match.group(1), int(match.group(2))) if match else ("X", 999)


def evidence_register(bundle):
    counters, register = {}, {}
    for item in sorted(bundle.get("evidence", []), key=lambda e: e.get("collectedAt", "")):
        prefix = EVIDENCE_PREFIX.get(item.get("kind"), "D")
        counters[prefix] = counters.get(prefix, 0) + 1
        register[item["id"]] = {**item, "label": f"{prefix}{counters[prefix]:02d}", "prefix": prefix}
    return register


def readiness_status(assessment, register):
    status = (assessment or {}).get("status", "unassessed")
    if status in ("minor", "major"):
        kinds = [register[e]["kind"] for e in assessment.get("evidenceIds", []) if e in register]
        if kinds and all(kind == "interview" for kind in kinds):
            return "Not evidenced"
    return {"conforming": "Met", "observation": "Met", "minor": "Partly met", "major": "Not met",
            "not_applicable": "Not applicable (provisional)"}.get(status, "Not sampled")


def priority_for(severity, critical):
    if severity == "major" or (severity == "minor" and critical):
        return "Critical"
    return {"minor": "Important", "observation": "Improvement"}.get(severity, "")


def criterion_rows(bundle, register):
    requirements = {r["id"]: r for r in bundle.get("requirements", [])}
    assessments = {a["requirementId"]: a for a in bundle.get("assessments", [])}
    findings = bundle.get("findings", [])
    actions = bundle.get("actions", [])
    result, custom = [], 0
    for requirement_id in bundle["audit"].get("requirementIds", []):
        r = requirements.get(requirement_id)
        a = assessments.get(requirement_id)
        if not r:
            continue
        code = (r.get("source") or {}).get("clause") or ""
        if not re.match(r"^[A-D]\d+$", code):
            custom += 1
            code = f"X{custom}"
        severity = (a or {}).get("status", "unassessed")
        linked = [register[e] for e in (a or {}).get("evidenceIds", []) if e in register]
        own_findings = [f for f in findings if f.get("requirementId") == requirement_id and f.get("status") != "withdrawn"]
        result.append({
            "id": requirement_id, "code": code, "pillar": code[0],
            "title": r.get("title") or r.get("text", "")[:120],
            "statement": r.get("text", ""),
            "indicators": r.get("indicators") or [],
            "critical": bool(r.get("critical")),
            "severity": severity,
            "status": readiness_status(a, register),
            "priority": priority_for(severity, bool(r.get("critical"))),
            "rationale": (a or {}).get("rationale", ""),
            "evidence": [{"label": e["label"], "kind": e["kind"], "description": e.get("description", "")} for e in linked],
            "findings": [f.get("statement", "") for f in own_findings],
            "actions": [x for x in actions if any(x.get("findingId") == f["id"] for f in own_findings) and x.get("status") != "cancelled"],
        })
    return sorted(result, key=lambda row: _clause_key(row["code"]))


def pillar_level(counts):
    applicable = sum(counts[s] for s in STATUSES[:4])
    if not applicable:
        return "Not reviewed"
    share = (counts["Met"] + 0.5 * counts["Partly met"]) / applicable
    level = "Advanced" if share >= 0.75 else "Developing" if share >= 0.4 else "Early"
    return f"{level} (incomplete)" if counts["Not sampled"] else level


def _text(value, limit):
    return value.strip()[:limit] if isinstance(value, str) and value.strip() else ""


def clean_narrative(raw, criteria):
    """Keep only well-formed AI output that refers to criteria in this audit."""
    raw = raw if isinstance(raw, dict) else {}
    codes = {row["code"] for row in criteria}
    pillars = {}
    for code, _ in PILLARS + [OTHER]:
        entry = (raw.get("pillars") or {}).get(code) if isinstance(raw.get("pillars"), dict) else None
        if isinstance(entry, dict):
            pillars[code] = {k: _text(entry.get(k), 600) for k in ("headline", "in_place", "missing")}
    per_criterion = {}
    if isinstance(raw.get("criteria"), dict):
        for code, entry in raw["criteria"].items():
            if code in codes and isinstance(entry, dict):
                per_criterion[code] = {"evidence_seen": _text(entry.get("evidence_seen"), 800), "gap": _text(entry.get("gap"), 1200)}
    strengths = [{"text": _text(s.get("text"), 500), "evidence": s.get("evidence") if s.get("evidence") in EVIDENCE_TYPES else "unverified"}
                 for s in raw.get("strengths") or [] if isinstance(s, dict) and _text(s.get("text"), 500)][:6]
    top_gaps = [{"criterion": g["criterion"], "text": _text(g.get("text"), 500)}
                for g in raw.get("top_gaps") or [] if isinstance(g, dict) and g.get("criterion") in codes and _text(g.get("text"), 500)][:5]
    actions = []
    for item in raw.get("actions") or []:
        if not isinstance(item, dict) or not _text(item.get("action"), 1000):
            continue
        linked = [c for c in (item.get("criteria") or []) if isinstance(c, str) and c in codes]
        if linked:
            actions.append({"criteria": list(dict.fromkeys(linked)), "action": _text(item.get("action"), 1000),
                            "owner": _text(item.get("owner"), 200) or PLACEHOLDER, "evidence": _text(item.get("evidence"), 800) or PLACEHOLDER})
    return {
        "letter_summary": _text(raw.get("letter_summary"), 2000),
        "readiness_statement": _text(raw.get("readiness_statement"), 1000),
        "limitations": _text(raw.get("limitations"), 1500),
        "pillars": pillars, "criteria": per_criterion, "strengths": strengths, "top_gaps": top_gaps, "actions": actions[:60],
    }


def ai_facts(bundle, profile):
    """The only audit data sent to the AI provider: text records, never files."""
    register = evidence_register(bundle)
    criteria = criterion_rows(bundle, register)
    return {
        "hotel": {k: profile.get(k, "") for k in ("hotel_name", "location", "hotel_type", "rooms", "staff", "facilities", "occupancy",
                                                   "water_sources", "wastewater", "existing_certifications", "review_date", "staff_present",
                                                   "areas_inspected", "areas_not_inspected", "plan_changes", "key_figures", "legal_items")},
        "standard": profile.get("standard_used", ""),
        "criteria": [{
            "code": row["code"], "title": row["title"], "status": row["status"], "priority": row["priority"] or None,
            "critical_flag": row["critical"], "auditor_rationale": row["rationale"],
            "evidence": [{"id": e["label"], "type": EVIDENCE_LABEL.get(e["kind"], e["kind"]), "description": e["description"]} for e in row["evidence"]],
            "findings": row["findings"],
            "assigned_actions": [x.get("description", "") for x in row["actions"]],
            **({"statement": row["statement"], "indicators": row["indicators"]} if row["status"] in NEEDS_WORK or row["priority"] == "Improvement" else {}),
        } for row in criteria],
    }


NOTICE = [
    "This is a limited, sample-based readiness assessment of the evidence made available on {date}. It is not certification, accreditation or a legal compliance determination, is not issued on behalf of any certification body, and does not guarantee the outcome of any certification audit.",
    "One day samples; it does not verify everything. Findings reflect the areas visited, documents seen and people interviewed on the day. A criterion rated “Met” means the evidence seen supported it, not that it is met at all times.",
    "The hotel remains responsible for its own legal compliance. Where this report mentions laws or local rules, it flags them for the hotel to confirm with the relevant authority or its own advisers.",
    "Standards change. Findings are against the GSTC Hotel Standard {edition} as published at the review date. Each certification scheme applies its own rules on top of it.",
    "Independence. AXIS does not certify hotels. AXIS is not affiliated with, or endorsed by, GSTC or any certification scheme.",
]
STATUS_MEANINGS = [
    ("Met", "The evidence seen supports every indicator reviewed"),
    ("Partly met", "Some indicators are met, or the practice exists but is not documented or consistent"),
    ("Not met", "The practice is absent, or evidence shows it is not in place"),
    ("Not evidenced", "The hotel says it happens, but no evidence was available on the day"),
    ("Not sampled", "Not reviewed on this visit; must be checked before an audit"),
    ("Not applicable (provisional)", "Appears not to apply; give the reason and confirm with the certification body"),
]
PRIORITY_MEANINGS = [
    ("Critical", "Likely to prevent certification. Close before booking an audit."),
    ("Important", "Likely to be raised by an auditor. Close before the audit if possible."),
    ("Improvement", "Good practice beyond what an auditor is likely to require."),
]
METHOD_NOTES = [
    "Each finding is also labelled by evidence type: observed (seen on site), documented (record or policy seen), staff-reported (said in interview, not verified) or unverified (claimed, no evidence). An auditor will expect documented or observed evidence for every criterion.",
    "Every item of evidence has an ID so each finding can be traced: D01… documents and records, I01… interviews, O01… observations and P01… photos (Appendices A and B). A verbal statement alone cannot show sustained practice, and an item not reviewed is not a failure.",
    "How statuses are derived in AXIS: Conforming and Observation assessments are shown as Met (observations become Improvement items); Minor findings as Partly met; Major findings as Not met; a Minor or Major finding supported only by interview evidence as Not evidenced; unassessed criteria as Not sampled. Major findings, and Minor findings on criteria flagged critical, are Critical priority.",
]
CERTIFICATION_WORKS = "Hotels are certified; certification bodies are accredited. GSTC writes the standard but does not certify hotels itself. A hotel is certified by an independent certification body that GSTC has accredited, after an audit against the GSTC Hotel Standard."
ROUTE_STEPS = [
    "Close the gaps in section 7 and keep records running.",
    "Shortlist certification bodies. GSTC publishes the list of accredited certification bodies on its website. Check which ones audit in the hotel's country.",
    "Request quotes. Compare total cost over the certificate’s life, including annual surveillance and auditor travel.",
    "Complete the certification body’s application and self-assessment, and send the documents it asks for.",
    "Certification audit. An on-site audit, usually followed by a period to correct any nonconformities before a decision.",
    "Keep it up. Maintain records and targets; expect surveillance or renewal audits.",
]
CB_QUESTIONS = [
    "Are you accredited by GSTC for hotels, and will you audit us against the Hotel Standard {edition}?",
    "Which additional GSTC or scheme requirements do you apply, and what thresholds and evidence periods do you expect?",
    "How many months of monitoring records do you expect at the first audit?",
    "What does certification cost over its full term, including surveillance audits and travel?",
    "Can the audit be run in the local language as well as English?",
    "How long do we have to correct findings raised at the audit?",
    "Do your rules restrict who may advise us before the audit?",
    "Which logos and claims may we use once certified?",
]
SCHEMES_NOTE = "GSTC-recognised schemes. Some schemes use their own standards that GSTC recognises as equivalent to its criteria. If the hotel’s goal is to be described as GSTC-certified, ask each scheme whether its certificate is issued by a GSTC-accredited certification body, and what claim it allows."
ACTION_CLOSURE = "An action counts as closed only when operating records show it is working, not when a document has been written. A policy on its own is a start; the auditor will also look for the practice it describes."


def build_readiness(bundle, saved_profile=None, draft=None, audit_version=None):
    profile = merge_profile(bundle, saved_profile)
    register = evidence_register(bundle)
    criteria = criterion_rows(bundle, register)
    narrative = clean_narrative((draft or {}).get("narrative"), criteria)
    has_ai = bool(draft and draft.get("narrative"))
    hotel = profile.get("hotel_name") or PLACEHOLDER
    edition = profile.get("standard_edition") or ""
    review_date = profile.get("review_date") or "[review date]"

    for row in criteria:
        ai = narrative["criteria"].get(row["code"], {})
        labels = ", ".join(e["label"] for e in row["evidence"])
        fallback_seen = f"{labels}: {row['rationale'][:220]}" if labels else (row["rationale"][:220] or "No evidence linked")
        seen = ai.get("evidence_seen")
        if seen and labels and not any(e["label"] in seen for e in row["evidence"]):
            seen = f"{labels}: {seen}"
        row["evidence_seen"] = seen or fallback_seen
        assigned = "; ".join(x.get("description", "") for x in row["actions"])
        row["gap"] = ai.get("gap") or assigned or (PLACEHOLDER if row["status"] in NEEDS_WORK or row["priority"] else "")

    pillars = []
    for code, name in PILLARS + [OTHER]:
        members = [row for row in criteria if row["pillar"] == code]
        if not members:
            continue
        counts = {status: sum(1 for row in members if row["status"] == status) for status in STATUSES}
        text = narrative["pillars"].get(code, {})
        pillars.append({
            "code": code, "name": name, "level": pillar_level(counts), "counts": counts, "total": len(members),
            "headline": text.get("headline") or PLACEHOLDER, "in_place": text.get("in_place") or "", "missing": text.get("missing") or "",
            "rows": [row for row in members if row["status"] in NEEDS_WORK or row["priority"]],
        })

    totals = {status: sum(p["counts"][status] for p in pillars) for status in STATUSES}
    needs = totals["Partly met"] + totals["Not met"] + totals["Not evidenced"]
    statement = (f"Of the {len(criteria)} criteria reviewed, {needs} need work before an audit: "
                 f"{totals['Partly met']} partly met, {totals['Not met']} not met and {totals['Not evidenced']} not evidenced.")
    if totals["Not sampled"]:
        statement += f" {totals['Not sampled']} were not sampled on this visit and must be checked before an audit."

    # Action plan: Critical-linked actions first (months 1–3), the rest in months 4–6.
    priority_of = {row["code"]: row["priority"] for row in criteria}
    proposed = narrative["actions"] or [
        {"criteria": [row["code"]], "action": row["gap"], "owner": PLACEHOLDER, "evidence": PLACEHOLDER}
        for row in criteria if row["status"] in NEEDS_WORK
    ]
    phase1 = [a for a in proposed if any(priority_of.get(c) == "Critical" for c in a["criteria"])]
    phase2 = [a for a in proposed if a not in phase1]
    numbered = [{**a, "number": i + 1} for i, a in enumerate(phase1 + phase2)]
    action_numbers = {}
    for action in numbered:
        for c in action["criteria"]:
            action_numbers.setdefault(c, []).append(action["number"])

    review_notes = []
    uncovered = [row["code"] for row in criteria if row["status"] in NEEDS_WORK and row["code"] not in action_numbers]
    if uncovered:
        review_notes.append(f"No action in section 7 covers: {', '.join(uncovered)}. Add actions before issue.")
    if not has_ai:
        review_notes.append("No AI narrative has been generated. Sections marked “[Reviewer to complete]” need writing before issue.")
    elif draft.get("audit_version") != audit_version:
        review_notes.append("The audit changed after the AI narrative was generated. Regenerate it, or check the narrative against the current assessments.")
    for field, label in (("hotel_contact_name", "hotel contact"), ("review_date", "review date"), ("report_reference", "report reference")):
        if not profile.get(field):
            review_notes.append(f"Add the {label} in Report details.")

    finding_code = {}
    for row in criteria:
        for f in bundle.get("findings", []):
            if f.get("requirementId") == row["id"]:
                finding_code[f["id"]] = row["code"]
    assigned_actions = [{
        "criterion": finding_code.get(x.get("findingId"), "—"), "description": x.get("description", ""),
        "owner": (x.get("owner") or {}).get("name", ""), "due": (x.get("dueDate") or "")[:10], "status": x.get("status", "").replace("_", " "),
    } for x in bundle.get("actions", []) if x.get("status") != "cancelled"]

    used_by = {}
    for row in criteria:
        for e in row["evidence"]:
            used_by.setdefault(e["label"], []).append(row["code"])
    evidence_rows = []
    for item in register.values():
        evidence_rows.append({
            "id": item["label"], "prefix": item["prefix"], "kind": EVIDENCE_LABEL.get(item.get("kind"), item.get("kind", "")),
            "description": item.get("description", ""),
            "reference": " · ".join(filter(None, [item.get("reference"), (item.get("attachment") or {}).get("fileName")])),
            "criteria": ", ".join(used_by.get(item["label"], [])) or "—",
            "date": (item.get("collectedAt") or "")[:10], "by": (item.get("collectedBy") or {}).get("name", ""),
        })

    contact = ", ".join(filter(None, [profile.get("hotel_contact_name"), profile.get("hotel_contact_role")])) or PLACEHOLDER
    ai_meta = None
    if has_ai:
        ai_meta = {k: draft.get(k) for k in ("model", "generated_at", "generated_by", "audit_version")}
        ai_meta["stale"] = draft.get("audit_version") != audit_version

    return {
        "draft": not profile.get("reviewed_by"),
        "ai": ai_meta,
        "profile": profile,
        "title": "AXIS Sustainability Readiness Review",
        "hotel": hotel,
        "cover": [
            ("Hotel", hotel), ("Location", profile.get("location") or PLACEHOLDER), ("Hotel contact", contact),
            ("Purpose", profile["purpose"]), ("Standard used", profile.get("standard_used") or PLACEHOLDER),
            ("Review date", review_date), ("Reviewer", profile.get("reviewer") or PLACEHOLDER),
            ("Report reference", profile.get("report_reference") or PLACEHOLDER), ("Report date", profile.get("report_date") or PLACEHOLDER),
            ("Confidentiality", profile["confidentiality"]),
        ],
        "letter": {
            "salutation": f"Dear {profile.get('hotel_contact_name') or '[Name]'},",
            "paragraphs": [
                f"Thank you for hosting the AXIS Sustainability Readiness Review at {hotel} on {review_date}"
                + (f", with {profile['staff_present']}." if profile.get("staff_present") else "."),
                f"You asked us to identify what {hotel} needs to address to achieve GSTC certification. This report compares what we saw on the day with "
                f"{'all ' if len(criteria) == 40 else ''}{len(criteria)} criteria of the {profile.get('standard_used') or 'GSTC Hotel Standard'}, lists every gap a certification auditor would be likely to raise, "
                "and sets out what to fix, in what order. It is a readiness review, not a certification audit, and it does not award a score or a pass.",
                narrative["letter_summary"] or PLACEHOLDER,
                "The certification action plan in section 7 lists what to fix first, and section 8 explains the route to certification. "
                "We recommend a follow-up review before you book an audit with a GSTC-accredited certification body.",
                "Thank you to your team for their time and openness.",
            ],
            "signoff": ["Kind regards,", profile.get("reviewer") or "[Name]", "AXIS", profile.get("reviewer_contact") or ""],
        },
        "notice": [p.format(date=review_date, edition=edition) for p in NOTICE],
        "summary": {
            "statement": statement, "readiness_statement": narrative["readiness_statement"],
            "pillars": pillars, "totals": totals, "strengths": narrative["strengths"],
            "top_gaps": narrative["top_gaps"] or [{"criterion": row["code"], "text": row["gap"]} for row in criteria if row["priority"] == "Critical"][:5],
            "limitations": narrative["limitations"],
        },
        "scope": {
            "certification_works": CERTIFICATION_WORKS,
            "table": [("Standard", profile.get("standard_used") or PLACEHOLDER), ("Site", ", ".join(filter(None, [hotel, profile.get("rooms") and f"{profile['rooms']} rooms", profile.get("location")]))),
                      ("Review type", profile["review_type"]), ("Duration", profile["duration"]), ("Method", profile["method"])],
            "people": rows(profile.get("people_interviewed"), 3),
            "areas_inspected": lines(profile.get("areas_inspected")),
            "areas_not_inspected": lines(profile.get("areas_not_inspected")),
            "status_meanings": STATUS_MEANINGS, "priority_meanings": PRIORITY_MEANINGS, "method_notes": METHOD_NOTES,
            "plan_changes": profile.get("plan_changes") or "No changes to the review plan were recorded.",
        },
        "hotel_profile": [(label, profile.get(key)) for key, label in (
            ("hotel_type", "Type"), ("rooms", "Rooms"), ("staff", "Staff"), ("facilities", "Facilities"), ("occupancy", "Occupancy, last 12 months"),
            ("water_sources", "Water sources"), ("wastewater", "Wastewater"), ("existing_certifications", "Existing certifications")) if profile.get(key)],
        "key_figures": rows(profile.get("key_figures"), 5),
        "gaps": pillars,
        "legal": rows(profile.get("legal_items"), 4),
        "plan": {"phase1": numbered[:len(phase1)], "phase2": numbered[len(phase1):], "closure_note": ACTION_CLOSURE},
        "assigned_actions": assigned_actions,
        "route": {"steps": ROUTE_STEPS, "questions": [q.format(edition=edition) for q in CB_QUESTIONS], "schemes": SCHEMES_NOTE},
        "evidence_register": [e for e in evidence_rows if e["prefix"] in ("D", "I")],
        "observations": [e for e in evidence_rows if e["prefix"] in ("O", "P")],
        "coverage": [{"code": row["code"], "title": row["title"], "status": row["status"],
                      "actions": ", ".join(map(str, action_numbers.get(row["code"], []))) or "—"} for row in criteria],
        "signoff": [("Prepared by", profile.get("prepared_by") or "[Name]"), ("Reviewed by", profile.get("reviewed_by") or "[Name]"),
                    ("Issued to", profile.get("issued_to") or contact), ("Next step", profile.get("next_step") or "Follow-up readiness review, [month]")],
        "review_notes": review_notes,
        "disclaimer": (bundle.get("template") or {}).get("disclaimer", ""),
    }


# ------------------------------------------------------------------ Markdown

def _cell(value):
    return str(value or "").replace("|", "\\|").replace("\n", " ")


def _table(headers, body):
    out = ["| " + " | ".join(headers) + " |", "|" + "---|" * len(headers)]
    out += ["| " + " | ".join(_cell(c) for c in row) + " |" for row in body]
    return out + [""]


def render_markdown(r):
    L = [f"# {r['title']}", ""]
    if r["draft"]:
        L += ["> **DRAFT — not reviewed.** " + ("Narrative sections were drafted with AI assistance from the audit record. " if r["ai"] else "")
              + "Check every statement against the evidence before issue.", ""]
    L += _table(["Cover details", ""], r["cover"])
    L += ["## Covering letter", "", r["letter"]["salutation"], ""] + [p + "\n" for p in r["letter"]["paragraphs"]] + ["  \n".join(filter(None, r["letter"]["signoff"])), ""]
    L += ["## 1. Important notice", ""] + [p + "\n" for p in r["notice"]]
    s = r["summary"]
    L += ["## 2. Summary of results", "", s["statement"] + (" " + s["readiness_statement"] if s["readiness_statement"] else ""), ""]
    L += _table(["Pillar", "Readiness", "Met", "Partly met", "Not met", "Not evidenced", "Not sampled", "Headline"],
                [[f"{p['code']}. {p['name']}", p["level"], p["counts"]["Met"], p["counts"]["Partly met"], p["counts"]["Not met"],
                  p["counts"]["Not evidenced"], p["counts"]["Not sampled"], p["headline"]] for p in s["pillars"]])
    if s["strengths"]:
        L += ["**Main strengths**", ""] + [f"- {x['text']} ({x['evidence']})" for x in s["strengths"]] + [""]
    if s["top_gaps"]:
        L += ["**Top gaps to close before a certification audit** (full list in section 7)", ""] + [f"{i + 1}. {g['text']} ({g['criterion']})" for i, g in enumerate(s["top_gaps"])] + [""]
    if s["limitations"]:
        L += [f"**What one day could not establish.** {s['limitations']}", ""]
    sc = r["scope"]
    L += ["## 3. Scope and method", "", "**How GSTC certification works.** " + sc["certification_works"], ""]
    L += _table(["Scope of this review", ""], sc["table"])
    if sc["people"]:
        L += ["**People interviewed**", ""] + _table(["Role", "Name", "Topics"], sc["people"])
    if sc["areas_inspected"] or sc["areas_not_inspected"]:
        L += ["**Areas inspected**", ""] + [f"- [x] {a}" for a in sc["areas_inspected"]] + [f"- [ ] {a} — not inspected" for a in sc["areas_not_inspected"]] + [""]
    L += ["**How findings are rated**", ""] + _table(["Status", "Meaning"], sc["status_meanings"]) + _table(["Priority", "Meaning for certification"], sc["priority_meanings"])
    L += [n + "\n" for n in sc["method_notes"]] + ["**Changes to the review plan.** " + sc["plan_changes"], ""]
    L += ["## 4. Hotel profile and key figures", ""]
    L += _table(["Hotel profile", ""], r["hotel_profile"]) if r["hotel_profile"] else [PLACEHOLDER, ""]
    if r["key_figures"]:
        L += _table(["Indicator (last 12 months)", "Total", "Per occupied room night", "Data status", "What certification needs"], r["key_figures"])
    L += ["## 5. Gap analysis: what to address for certification", "",
          "For each criterion, the “Gap to close” column says what must exist on the day of a certification audit. Gaps marked Critical are the ones most likely to stop certification. Appendix C records a status for every criterion.", ""]
    for p in r["gaps"]:
        L += [f"### {p['code']}. {p['name']}", ""]
        if p["in_place"] or p["missing"]:
            L += [" ".join(filter(None, [p["in_place"] and f"In place: {p['in_place']}", p["missing"] and f"Missing: {p['missing']}"])), ""]
        L += _table(["Code", "Criterion", "Status", "Evidence seen", "Gap to close", "Priority"],
                    [[x["code"], x["title"], x["status"], x["evidence_seen"], x["gap"], x["priority"] or "—"] for x in p["rows"]]) if p["rows"] else ["No gaps recorded in this pillar.", ""]
    L += ["## 6. Local legal requirements an auditor will check", "",
          "GSTC criterion A2 requires the hotel to know and comply with every law that applies to it. AXIS does not make legal compliance determinations: the hotel should confirm each item with the relevant authority or its own adviser.", ""]
    L += _table(["Area", "What to confirm", "What we saw", "Priority"], r["legal"]) if r["legal"] else [PLACEHOLDER + " — add local legal items in Report details.", ""]
    L += ["## 7. Certification action plan", "", "Close all Critical gaps first, then build up the records an auditor will want to see. Ask your chosen certification body how many months of records it expects.", ""]
    for title, items in (("Months 1–3: close the Critical gaps", r["plan"]["phase1"]), ("Months 4–6: build the evidence", r["plan"]["phase2"])):
        L += [f"**{title}**", ""] + (_table(["#", "Action", "Criterion", "Owner", "Evidence the auditor will want"],
                                            [[a["number"], a["action"], ", ".join(a["criteria"]), a["owner"], a["evidence"]] for a in items]) if items else ["None.", ""])
    L += [r["plan"]["closure_note"], ""]
    if r["assigned_actions"]:
        L += ["**Corrective actions already assigned in AXIS**", ""] + _table(["Criterion", "Action", "Owner", "Due", "Status"],
                                                                            [[a["criterion"], a["description"], a["owner"], a["due"], a["status"]] for a in r["assigned_actions"]])
    L += ["**Month 6: check readiness, then book**", "", "- [ ] AXIS follow-up review against the same criteria", "- [ ] Choose a GSTC-accredited certification body and request a quote (section 8)", "- [ ] Book the certification audit once no Critical gaps remain", ""]
    L += ["## 8. Route to GSTC certification", "", "Once the Critical gaps are closed, the hotel chooses a GSTC-accredited certification body, which audits it and issues the certificate. AXIS cannot certify the hotel and has no commercial link with any certification body.", ""]
    L += [f"{i + 1}. {step}" for i, step in enumerate(r["route"]["steps"])] + ["", "**Questions to ask each certification body**", ""] + [f"- [ ] {q}" for q in r["route"]["questions"]] + ["", r["route"]["schemes"], ""]
    L += ["## Appendix A: Evidence register", ""] + (_table(["ID", "Evidence", "Criterion", "Type", "Reference", "Date"], [[e["id"], e["description"], e["criteria"], e["kind"], e["reference"], e["date"]] for e in r["evidence_register"]]) if r["evidence_register"] else ["No documents or interviews recorded.", ""])
    L += ["## Appendix B: Observation and photo log", ""] + (_table(["ID", "What was observed", "Criterion", "Type", "Reference"], [[e["id"], e["description"], e["criteria"], e["kind"], e["reference"]] for e in r["observations"]]) if r["observations"] else ["No observations or photos recorded.", ""])
    L += ["## Appendix C: Criterion coverage register", ""] + _table(["Code", "Criterion", "Status", "Action (section 7)"], [[c["code"], c["title"], c["status"], c["actions"]] for c in r["coverage"]])
    L += ["## Sign-off", ""] + _table(["", ""], r["signoff"])
    if r["disclaimer"]:
        L += [f"*{r['disclaimer']}*", ""]
    return "\n".join(L)


# ------------------------------------------------------------------ Word

def render_docx(r):
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt, RGBColor

    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)

    def shade(cell, fill):
        props = cell._tc.get_or_add_tcPr()
        element = OxmlElement("w:shd")
        element.set(qn("w:val"), "clear")
        element.set(qn("w:color"), "auto")
        element.set(qn("w:fill"), fill)
        props.append(element)

    def table(headers, body, widths=None):
        t = doc.add_table(rows=1, cols=len(headers))
        t.style = "Table Grid"
        for i, h in enumerate(headers):
            cell = t.rows[0].cells[i]
            cell.text = ""
            run = cell.paragraphs[0].add_run(str(h))
            run.bold = True
            run.font.size = Pt(9.5)
            shade(cell, "DCE6F0")
        for row in body:
            cells = t.add_row().cells
            for i, value in enumerate(row):
                cells[i].text = ""
                run = cells[i].paragraphs[0].add_run(str(value if value is not None else ""))
                run.font.size = Pt(9.5)
        doc.add_paragraph()
        return t

    def pairs(body):
        t = doc.add_table(rows=0, cols=2)
        t.style = "Table Grid"
        for label, value in body:
            cells = t.add_row().cells
            cells[0].text = ""
            cells[0].paragraphs[0].add_run(str(label)).bold = True
            shade(cells[0], "F2F2F2")
            cells[1].text = str(value or "")
        doc.add_paragraph()

    def para(text, bold_lead=None, italic=False):
        p = doc.add_paragraph()
        if bold_lead:
            p.add_run(bold_lead + " ").bold = True
        run = p.add_run(text)
        run.italic = italic
        return p

    title = doc.add_heading(r["title"], 0)
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    if r["draft"]:
        warn = doc.add_paragraph()
        run = warn.add_run("DRAFT — not reviewed. " + ("Narrative sections were drafted with AI assistance from the audit record. " if r["ai"] else "")
                           + "Check every statement against the evidence, then enter a reviewer in Report details before issue.")
        run.bold = True
        run.font.color.rgb = RGBColor(0xB0, 0x3A, 0x2E)
    doc.add_heading("Cover details", 1)
    pairs(r["cover"])

    doc.add_heading("Covering letter", 1)
    para(r["letter"]["salutation"])
    for p in r["letter"]["paragraphs"]:
        para(p)
    for line in filter(None, r["letter"]["signoff"]):
        doc.add_paragraph(line)

    doc.add_heading("1. Important notice", 1)
    for p in r["notice"]:
        para(p)

    s = r["summary"]
    doc.add_heading("2. Summary of results", 1)
    para(s["statement"] + (" " + s["readiness_statement"] if s["readiness_statement"] else ""))
    table(["Pillar", "Readiness", "Met", "Partly met", "Not met", "Not evidenced", "Not sampled", "Headline"],
          [[f"{p['code']}. {p['name']}", p["level"], p["counts"]["Met"], p["counts"]["Partly met"], p["counts"]["Not met"],
            p["counts"]["Not evidenced"], p["counts"]["Not sampled"], p["headline"]] for p in s["pillars"]])
    if s["strengths"]:
        doc.add_heading("Main strengths", 3)
        for x in s["strengths"]:
            doc.add_paragraph(f"{x['text']} ({x['evidence']})", style="List Bullet")
    if s["top_gaps"]:
        doc.add_heading("Top gaps to close before a certification audit", 3)
        for g in s["top_gaps"]:
            doc.add_paragraph(f"{g['text']} ({g['criterion']})", style="List Number")
    if s["limitations"]:
        para(s["limitations"], "What one day could not establish.")

    sc = r["scope"]
    doc.add_heading("3. Scope and method", 1)
    para(sc["certification_works"], "How GSTC certification works.")
    doc.add_heading("Scope of this review", 3)
    pairs(sc["table"])
    if sc["people"]:
        doc.add_heading("People interviewed", 3)
        table(["Role", "Name", "Topics"], sc["people"])
    if sc["areas_inspected"] or sc["areas_not_inspected"]:
        doc.add_heading("Areas inspected", 3)
        for a in sc["areas_inspected"]:
            doc.add_paragraph(f"☒ {a}")
        for a in sc["areas_not_inspected"]:
            doc.add_paragraph(f"☐ {a} — not inspected")
    doc.add_heading("How findings are rated", 3)
    table(["Status", "Meaning"], sc["status_meanings"])
    table(["Priority", "Meaning for certification"], sc["priority_meanings"])
    for n in sc["method_notes"]:
        para(n)
    para(sc["plan_changes"], "Changes to the review plan.")

    doc.add_heading("4. Hotel profile and key figures", 1)
    if r["hotel_profile"]:
        pairs(r["hotel_profile"])
    else:
        para(PLACEHOLDER)
    if r["key_figures"]:
        table(["Indicator (last 12 months)", "Total", "Per occupied room night", "Data status", "What certification needs"], r["key_figures"])

    doc.add_heading("5. Gap analysis: what to address for certification", 1)
    para("For each criterion, the “Gap to close” column says what must exist on the day of a certification audit. Gaps marked Critical are the ones most likely to stop certification. Appendix C records a status for every criterion.")
    for p in r["gaps"]:
        doc.add_heading(f"{p['code']}. {p['name']}", 2)
        if p["in_place"]:
            para(p["in_place"], "In place:")
        if p["missing"]:
            para(p["missing"], "Missing:")
        if p["rows"]:
            table(["Code", "Criterion", "Status", "Evidence seen", "Gap to close", "Priority"],
                  [[x["code"], x["title"], x["status"], x["evidence_seen"], x["gap"], x["priority"] or "—"] for x in p["rows"]])
        else:
            para("No gaps recorded in this pillar.")

    doc.add_heading("6. Local legal requirements an auditor will check", 1)
    para("GSTC criterion A2 requires the hotel to know and comply with every law that applies to it. AXIS does not make legal compliance determinations: the hotel should confirm each item with the relevant authority or its own adviser.")
    if r["legal"]:
        table(["Area", "What to confirm", "What we saw", "Priority"], r["legal"])
    else:
        para(PLACEHOLDER + " — add local legal items in Report details.")

    doc.add_heading("7. Certification action plan", 1)
    para("Close all Critical gaps first, then build up the records an auditor will want to see. Ask your chosen certification body how many months of records it expects.")
    for heading, items in (("Months 1–3: close the Critical gaps", r["plan"]["phase1"]), ("Months 4–6: build the evidence", r["plan"]["phase2"])):
        doc.add_heading(heading, 3)
        if items:
            table(["#", "Action", "Criterion", "Owner", "Evidence the auditor will want"],
                  [[a["number"], a["action"], ", ".join(a["criteria"]), a["owner"], a["evidence"]] for a in items])
        else:
            para("None.")
    para(r["plan"]["closure_note"], italic=True)
    if r["assigned_actions"]:
        doc.add_heading("Corrective actions already assigned in AXIS", 3)
        table(["Criterion", "Action", "Owner", "Due", "Status"], [[a["criterion"], a["description"], a["owner"], a["due"], a["status"]] for a in r["assigned_actions"]])
    doc.add_heading("Month 6: check readiness, then book", 3)
    for item in ("AXIS follow-up review against the same criteria", "Choose a GSTC-accredited certification body and request a quote (section 8)", "Book the certification audit once no Critical gaps remain"):
        doc.add_paragraph(f"☐ {item}")

    doc.add_heading("8. Route to GSTC certification", 1)
    para("Once the Critical gaps are closed, the hotel chooses a GSTC-accredited certification body, which audits it and issues the certificate. AXIS cannot certify the hotel and has no commercial link with any certification body.")
    for step in r["route"]["steps"]:
        doc.add_paragraph(step, style="List Number")
    doc.add_heading("Questions to ask each certification body", 3)
    for q in r["route"]["questions"]:
        doc.add_paragraph(f"☐ {q}")
    para(r["route"]["schemes"])

    doc.add_heading("Appendix A: Evidence register", 1)
    if r["evidence_register"]:
        table(["ID", "Evidence", "Criterion", "Type", "Reference", "Date"], [[e["id"], e["description"], e["criteria"], e["kind"], e["reference"], e["date"]] for e in r["evidence_register"]])
    else:
        para("No documents or interviews recorded.")
    doc.add_heading("Appendix B: Observation and photo log", 1)
    if r["observations"]:
        table(["ID", "What was observed", "Criterion", "Type", "Reference"], [[e["id"], e["description"], e["criteria"], e["kind"], e["reference"]] for e in r["observations"]])
    else:
        para("No observations or photos recorded.")
    doc.add_heading("Appendix C: Criterion coverage register", 1)
    table(["Code", "Criterion", "Status", "Action (section 7)"], [[c["code"], c["title"], c["status"], c["actions"]] for c in r["coverage"]])

    doc.add_heading("Sign-off", 1)
    pairs(r["signoff"])
    if r["disclaimer"]:
        para(r["disclaimer"], italic=True)

    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def slug(value):
    return re.sub(r"[^A-Za-z0-9]+", "-", value or "hotel").strip("-")[:60] or "hotel"


def now_iso():
    return datetime.now(timezone.utc).replace(microsecond=0).isoformat()
